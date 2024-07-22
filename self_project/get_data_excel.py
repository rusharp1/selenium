import openpyxl

# 호텔 검색 페이지에 사용할 data값.
def search_get_excel_data(file_path):
    workbook = openpyxl.load_workbook(file_path)
    ws = workbook.active
    # 
    excel_value = []
    for row in ws.iter_rows(min_row = 2, max_col=6, values_only=True):
        row_list = list(row)
        if row[-1]:
            tmp = [item.strip() for item in str(row[-1]).split(',')]
        else:
            tmp = []
        row_list[-1] = tmp
        excel_value.append(row_list)
    # 
    return excel_value


#  호텔 검색 완료 페이지이서 사용할 필터 값.
def filter_get_excel_data(file_path):
    workbook = openpyxl.load_workbook(file_path)
    ws = workbook.active
    # 
    excel_value = []
    for row in ws.iter_rows(min_row = 2, min_col=7,max_col=13, values_only=True):
        row_list = list(row)
        for idx, row_value in enumerate(row):
            if row_value and ',' in str(row_value):
                tmp = [item.strip() for item in row_value.split(',')]
                row_list[idx] = tmp
            elif not(row_value):
                row_list[idx] = []
            else:
                row_list[idx] = [row_list[idx]]
        excel_value.append(row_list)
    # 
    return excel_value

if __name__== "__main__":
    tmp = filter_get_excel_data('data_value.xlsx')
    for v in tmp:
        print(v)
    tmp = search_get_excel_data('data_value.xlsx')
    for v in tmp:
        print(v)